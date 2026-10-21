"""Import a product catalogue from a user-supplied CSV / Excel file.

This NEVER logs in to Shopee, never uses cookies, and never scrapes. It only
reads a file the user exported themselves (from the Shopee Affiliate site's own
export, or a spreadsheet they typed). Prices, commission and links are taken as
provided and marked unverified — they are the user's data, not proof from an API.
"""
from __future__ import annotations

import csv
import io
import re
from urllib.parse import urlsplit

MAX_ROWS = 500
MAX_FILE_BYTES = 4 * 1024 * 1024

# Columns the importer understands, each with the normalised header aliases
# (Thai + English) that map onto it. normalize() strips spaces/punctuation/case.
FIELD_ALIASES = {
    'title': ['ชื่อสินค้า', 'ชื่อสินค้าที่สั่งซื้อ', 'ชื่อผลิตภัณฑ์', 'ชื่อรายการสินค้า', 'สินค้า',
              'productname', 'itemname', 'producttitle', 'itemtitle', 'name', 'title', 'product'],
    'price': ['ราคา', 'ราคาสินค้า', 'price', 'productprice', 'itemprice', 'sellingprice'],
    'aff_link': ['ลิงก์affiliate', 'ลิงค์affiliate', 'ลิงก์aff', 'ลิงค์aff', 'affiliatelink', 'afflink',
                 'shortlink', 'ลิงก์สินค้า', 'ลิงค์สินค้า', 'producturl', 'productlink', 'itemlink', 'link', 'url'],
    'image_url': ['urlภาพสินค้า', 'urlภาพ', 'ภาพสินค้า', 'รูปสินค้า', 'รูปภาพ', 'imageurl', 'image', 'img', 'imageurlสินค้า', 'photo', 'thumbnail'],
    'item_id': ['รหัสสินค้า', 'รหัส', 'itemid', 'productid', 'id', 'sku', 'itemcode'],
    'commission_rate': ['อัตราค่าคอมมิชชัน', 'อัตราค่าคอม', 'ค่าคอมเปอร์เซ็นต์', 'commissionrate', 'commission', 'commpct', 'rate'],
}
# Image hosts we will fetch server-side (Shopee product CDN only).
IMAGE_HOSTS = ('susercontent.com', 'shopee.co.th')
AFF_HOSTS = {'shopee.co.th', 'www.shopee.co.th', 's.shopee.co.th', 'shp.ee'}


def normalize(h) -> str:
    return re.sub(r'[\s_\-()\[\]฿%./]+', '', str(h).lower().strip())


def _https_host(raw: str, allowed, allow_suffix=False):
    """Return the url when it is https and its host is allowed, else ''."""
    try:
        raw = str(raw).strip().rstrip('),.;]}>')
        if not raw or len(raw) > 1000 or any(ch.isspace() for ch in raw):
            return ''
        p = urlsplit(raw)
        host = (p.hostname or '').lower()
        if p.scheme != 'https' or p.username or p.password or p.port not in (None, 443):
            return ''
        if allow_suffix:
            if not any(host == h or host.endswith('.' + h) for h in allowed):
                return ''
        elif host not in allowed:
            return ''
        return raw
    except ValueError:
        return ''


def aff_link(raw: str) -> str:
    return _https_host(raw, AFF_HOSTS, allow_suffix=True)


def image_link(raw: str) -> str:
    return _https_host(raw, IMAGE_HOSTS, allow_suffix=True)


def _num(raw, lo, hi):
    x = str(raw or '').replace(',', '').replace('฿', '').strip().rstrip('%').strip()
    if not x:
        return None
    try:
        v = float(x)
    except ValueError:
        return None
    return v if lo <= v <= hi else None


def _read_rows(binary: bytes, filename: str):
    """Return (columns, list-of-dict-rows) from a .csv/.txt or .xlsx file."""
    if len(binary) > MAX_FILE_BYTES:
        raise ValueError(f'ไฟล์เกิน {MAX_FILE_BYTES // (1024 * 1024)} MB')
    name = (filename or '').lower()
    if name.endswith('.xlsx'):
        try:
            import openpyxl
        except ImportError:
            raise ValueError('เซิร์ฟเวอร์ยังไม่รองรับ .xlsx — กรุณาบันทึกเป็น .csv')
        try:
            wb = openpyxl.load_workbook(io.BytesIO(binary), read_only=True, data_only=True)
        except Exception:
            raise ValueError('เปิดไฟล์ Excel ไม่ได้ — ไฟล์อาจเสียหาย')
        ws = wb[wb.sheetnames[0]]
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if not header:
            raise ValueError('ชีทแรกของไฟล์ว่างเปล่า')
        cols = [str(c).strip() if c is not None else f'col{i}' for i, c in enumerate(header)]
        out = []
        for r in rows_iter:
            if len(out) >= MAX_ROWS:
                break
            out.append({cols[i]: ('' if v is None else str(v).strip()) for i, v in enumerate(r) if i < len(cols)})
        wb.close()
        return cols, out
    if not name.endswith(('.csv', '.txt')):
        raise ValueError('รองรับไฟล์ .csv และ .xlsx เท่านั้น')
    text = None
    for enc in ('utf-8-sig', 'utf-16', 'cp874'):
        try:
            text = binary.decode(enc)
            if '\x00' in text[:80]:
                text = None
                continue
            break
        except (UnicodeDecodeError, UnicodeError):
            pass
    if text is None:
        raise ValueError('อ่านรหัสอักษรของไฟล์ไม่ได้ — บันทึกเป็น CSV UTF-8')
    try:
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=',;\t')
    except csv.Error:
        dialect = 'excel'
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    if not reader.fieldnames:
        raise ValueError('ไฟล์ไม่มีหัวคอลัมน์')
    cols = [c for c in reader.fieldnames if c]
    rows = []
    for r in reader:
        if len(rows) >= MAX_ROWS:
            break
        rows.append({k: ('' if v is None else str(v).strip()) for k, v in r.items() if k})
    return cols, rows


def guess_mapping(columns):
    """Best-guess {field: column} from the header names."""
    norm = {normalize(c): c for c in columns}
    mapping = {}
    for field, aliases in FIELD_ALIASES.items():
        for a in aliases:
            na = normalize(a)
            # exact first, then contains
            if na in norm:
                mapping[field] = norm[na]
                break
        else:
            for nc, c in norm.items():
                if any(normalize(a) in nc for a in aliases):
                    mapping[field] = c
                    break
    return mapping


def preview(binary: bytes, filename: str):
    cols, rows = _read_rows(binary, filename)
    if not rows:
        raise ValueError('ไม่พบข้อมูลสินค้าในไฟล์')
    return {
        'columns': cols[:40],
        'guess': guess_mapping(cols),
        'fields': list(FIELD_ALIASES.keys()),
        'sample': rows[:5],
        'total': len(rows),
    }


def build_items(binary: bytes, filename: str, mapping: dict):
    """Turn file rows into clean items using the user-confirmed column mapping."""
    cols, rows = _read_rows(binary, filename)
    colset = set(cols)
    m = {f: c for f, c in (mapping or {}).items() if c in colset}
    title_col = m.get('title')
    if not title_col:
        raise ValueError('กรุณาเลือกคอลัมน์ "ชื่อสินค้า" ก่อนนำเข้า')
    items, skipped = [], 0
    for r in rows:
        title = (r.get(title_col) or '').strip()[:200]
        if not title or title.lower() in ('total', 'รวม', 'sum'):
            skipped += 1
            continue
        item = {
            'title': title,
            'price': _num(r.get(m['price']), 0, 10_000_000) if 'price' in m else None,
            'commission_rate': _num(r.get(m['commission_rate']), 0, 100) if 'commission_rate' in m else None,
            'url': aff_link(r.get(m['aff_link'], '')) if 'aff_link' in m else '',
            'image_url': image_link(r.get(m['image_url'], '')) if 'image_url' in m else '',
            'item_id': re.sub(r'\D', '', str(r.get(m['item_id'], '')))[:20] if 'item_id' in m else '',
        }
        items.append(item)
    if not items:
        raise ValueError('ไม่พบสินค้าที่นำเข้าได้ — ตรวจสอบว่าเลือกคอลัมน์ชื่อสินค้าถูกต้อง')
    return items, skipped
